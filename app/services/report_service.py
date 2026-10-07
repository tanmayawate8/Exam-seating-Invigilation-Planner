"""
Report Service & Member 4 Gateway Module.
Assembles standard domain payloads according to Data Contracts and invokes
Member 4's ReportLab (PDF) and openpyxl (Excel) rendering engines.
"""

from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.contracts.report_contract import (
    AttendanceCandidateItem,
    AttendanceSheetReportContract,
    DutyRosterItem,
    DutyRosterReportContract,
    RoomNoticeCandidateItem,
    RoomNoticeReportContract,
    SeatingChartItem,
    SeatingChartReportContract,
)
from app.extensions import db
from app.models.exam import Exam
from app.models.invigilation_duty import InvigilationDuty
from app.models.room import Room
from app.models.seat_allocation import SeatAllocation
from app.models.seating_plan import SeatingPlan
from app.utils.audit import log_audit
from app.utils.errors import BadRequestError, NotFoundError

INSTITUTION_NAME = "POLYTECHNIC EXAMINATION CELL"


class ReportService:
    """Backend service gateway for generating institutional examination reports."""

    # =========================================================================
    # CONTRACT BUILDERS (Extracts normalized data from PostgreSQL)
    # =========================================================================

    @staticmethod
    def build_seating_chart_contract(plan_id: int) -> SeatingChartReportContract:
        """Assembles data contract payload for a master seating plan chart."""
        plan = SeatingPlan.query.get(plan_id)
        if not plan:
            raise NotFoundError(f"Seating Plan with ID {plan_id} not found.")

        exam = plan.exam
        allocations = (
            SeatAllocation.query.filter_by(seating_plan_id=plan.id)
            .join(SeatAllocation.student)
            .order_by(SeatAllocation.room_id, SeatAllocation.seat_id)
            .all()
        )

        items = []
        for alloc in allocations:
            st = alloc.student
            room = alloc.room
            seat = alloc.seat
            items.append(
                SeatingChartItem(
                    student_roll=st.roll_number,
                    student_name=st.full_name,
                    department_code=st.department.code if st.department else "GEN",
                    semester=st.semester,
                    room_number=room.room_number,
                    seat_number=seat.seat_number,
                    row_num=seat.row_num,
                    col_num=seat.col_num,
                )
            )

        return SeatingChartReportContract(
            institution_name=INSTITUTION_NAME,
            plan_code=plan.plan_code,
            exam_code=exam.exam_code,
            subject_code=exam.subject.code if exam.subject else "N/A",
            subject_name=exam.subject.name if exam.subject else exam.title,
            exam_date=exam.exam_date.isoformat(),
            start_time=exam.start_time.isoformat(),
            end_time=exam.end_time.isoformat(),
            session_name=exam.session_name,
            allocations=items,
        )

    @staticmethod
    def build_room_notice_contract(plan_id: int, room_id: int) -> RoomNoticeReportContract:
        """Assembles data contract payload for a specific room's door notice."""
        plan = SeatingPlan.query.get(plan_id)
        if not plan:
            raise NotFoundError(f"Seating Plan with ID {plan_id} not found.")

        room = Room.query.get(room_id)
        if not room:
            raise NotFoundError(f"Room with ID {room_id} not found.")

        exam = plan.exam
        allocations = (
            SeatAllocation.query.filter_by(seating_plan_id=plan.id, room_id=room.id)
            .join(SeatAllocation.seat)
            .order_by(SeatAllocation.seat_id)
            .all()
        )

        candidates = []
        for alloc in allocations:
            st = alloc.student
            candidates.append(
                RoomNoticeCandidateItem(
                    seat_number=alloc.seat.seat_number,
                    roll_number=st.roll_number,
                    student_name=st.full_name,
                    department_code=st.department.code if st.department else "GEN",
                    semester=st.semester,
                )
            )

        return RoomNoticeReportContract(
            institution_name=INSTITUTION_NAME,
            exam_code=exam.exam_code,
            subject_name=exam.subject.name if exam.subject else exam.title,
            exam_date=exam.exam_date.isoformat(),
            session_name=exam.session_name,
            room_number=room.room_number,
            building=room.building,
            floor=room.floor,
            capacity=room.capacity,
            candidates=candidates,
        )

    @staticmethod
    def build_attendance_sheet_contract(plan_id: int, room_id: int) -> AttendanceSheetReportContract:
        """Assembles data contract payload for a hall attendance signature sheet."""
        plan = SeatingPlan.query.get(plan_id)
        if not plan:
            raise NotFoundError(f"Seating Plan with ID {plan_id} not found.")

        room = Room.query.get(room_id)
        if not room:
            raise NotFoundError(f"Room with ID {room_id} not found.")

        exam = plan.exam
        allocations = (
            SeatAllocation.query.filter_by(seating_plan_id=plan.id, room_id=room.id)
            .join(SeatAllocation.student)
            .order_by(SeatAllocation.seat_id)
            .all()
        )

        # Check assigned invigilator
        duty = InvigilationDuty.query.filter_by(exam_id=exam.id, room_id=room.id).first()
        invigilator_name = duty.teacher.full_name if duty and duty.teacher else "Unassigned"

        candidates = []
        for idx, alloc in enumerate(allocations, start=1):
            st = alloc.student
            candidates.append(
                AttendanceCandidateItem(
                    sr_no=idx,
                    roll_number=st.roll_number,
                    enrollment_number=st.enrollment_number,
                    student_name=st.full_name,
                    seat_number=alloc.seat.seat_number,
                )
            )

        return AttendanceSheetReportContract(
            institution_name=INSTITUTION_NAME,
            exam_code=exam.exam_code,
            subject_code=exam.subject.code if exam.subject else "N/A",
            subject_name=exam.subject.name if exam.subject else exam.title,
            exam_date=exam.exam_date.isoformat(),
            session_name=exam.session_name,
            room_number=room.room_number,
            invigilator_name=invigilator_name,
            candidates=candidates,
        )

    @staticmethod
    def build_duty_roster_contract(exam_id: int) -> DutyRosterReportContract:
        """Assembles data contract payload for an exam session's faculty duty roster."""
        exam = Exam.query.get(exam_id)
        if not exam:
            raise NotFoundError(f"Exam with ID {exam_id} not found.")

        duties = (
            InvigilationDuty.query.filter_by(exam_id=exam.id)
            .join(InvigilationDuty.teacher)
            .order_by(InvigilationDuty.room_id)
            .all()
        )

        items = []
        for idx, d in enumerate(duties, start=1):
            t = d.teacher
            r = d.room
            items.append(
                DutyRosterItem(
                    sr_no=idx,
                    faculty_name=t.full_name,
                    employee_id=t.employee_id,
                    department_code=t.department.code if t.department else "GEN",
                    designation=t.designation,
                    room_number=r.room_number if r else "TBD",
                    building=r.building if r else "Main",
                    duty_role=d.duty_role,
                )
            )

        return DutyRosterReportContract(
            institution_name=INSTITUTION_NAME,
            exam_code=exam.exam_code,
            exam_title=exam.title,
            exam_date=exam.exam_date.isoformat(),
            session_name=exam.session_name,
            start_time=exam.start_time.isoformat(),
            end_time=exam.end_time.isoformat(),
            duties=items,
        )

    # =========================================================================
    # PDF RENDERING ENGINES (ReportLab)
    # =========================================================================

    @staticmethod
    def render_seating_chart_pdf(contract: SeatingChartReportContract) -> BytesIO:
        """Generates a professional PDF Master Seating Chart."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=15,
            leading=18,
            alignment=1,
            textColor=colors.HexColor("#1e3a8a"),
        )
        sub_style = ParagraphStyle(
            "SubStyle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            alignment=1,
            textColor=colors.HexColor("#475569"),
        )

        elements = [
            Paragraph(contract.institution_name, title_style),
            Paragraph(f"MASTER SEATING ARRANGEMENT CHART — {contract.plan_code}", sub_style),
            Paragraph(
                f"Exam: {contract.exam_code} | Subject: {contract.subject_name} ({contract.subject_code}) | "
                f"Date: {contract.exam_date} ({contract.session_name})",
                sub_style,
            ),
            Spacer(1, 15),
        ]

        # Table Header & Rows
        table_data = [["Sr", "Roll No", "Student Name", "Dept", "Sem", "Room", "Seat"]]
        for idx, item in enumerate(contract.allocations, start=1):
            table_data.append([
                str(idx),
                item.student_roll,
                item.student_name,
                item.department_code,
                f"Sem {item.semester}",
                item.room_number,
                item.seat_number,
            ])

        t = Table(table_data, colWidths=[30, 75, 175, 55, 45, 65, 75])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (2, 1), (2, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))

        elements.append(t)
        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def render_room_notice_pdf(contract: RoomNoticeReportContract) -> BytesIO:
        """Generates a Room Door Notice / Noticeboard Seating Chart PDF."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=18, leading=22, alignment=1)
        room_style = ParagraphStyle(
            "RoomHeader",
            parent=styles["Heading2"],
            fontSize=16,
            leading=20,
            alignment=1,
            textColor=colors.HexColor("#b91c1c"),
        )
        info_style = ParagraphStyle("Info", parent=styles["Normal"], fontSize=10, leading=14, alignment=1)

        elements = [
            Paragraph(contract.institution_name, title_style),
            Spacer(1, 4),
            Paragraph(f"EXAMINATION ROOM NOTICE: {contract.room_number}", room_style),
            Paragraph(f"Building: {contract.building} | Floor: {contract.floor} | Capacity: {contract.capacity}", info_style),
            Paragraph(f"Exam: {contract.exam_code} — {contract.subject_name}", info_style),
            Paragraph(f"Date: {contract.exam_date} | Session: {contract.session_name}", info_style),
            Spacer(1, 15),
        ]

        table_data = [["Seat No", "Roll Number", "Student Name", "Dept", "Sem"]]
        for c in contract.candidates:
            table_data.append([
                c.seat_number,
                c.roll_number,
                c.student_name,
                c.department_code,
                f"Sem {c.semester}",
            ])

        t = Table(table_data, colWidths=[90, 95, 200, 60, 60])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (2, 1), (2, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ]))

        elements.append(t)
        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def render_attendance_sheet_pdf(contract: AttendanceSheetReportContract) -> BytesIO:
        """Generates an Examination Attendance & Candidate Signature Sheet PDF."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("H1", parent=styles["Heading2"], fontSize=14, leading=17, alignment=1)
        sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=9, leading=12, alignment=1)

        elements = [
            Paragraph(contract.institution_name, title_style),
            Paragraph("EXAMINATION ATTENDANCE & HALL SIGNATURE SHEET", title_style),
            Paragraph(
                f"Exam: {contract.exam_code} | Subject: {contract.subject_name} ({contract.subject_code}) | "
                f"Room: {contract.room_number} | Invigilator: {contract.invigilator_name}",
                sub_style,
            ),
            Paragraph(f"Date: {contract.exam_date} | Session: {contract.session_name}", sub_style),
            Spacer(1, 10),
        ]

        table_data = [["Sr", "Roll No", "Enrollment No", "Student Name", "Seat", "Answer Book No", "Student Signature"]]
        for c in contract.candidates:
            table_data.append([
                str(c.sr_no),
                c.roll_number,
                c.enrollment_number,
                c.student_name,
                c.seat_number,
                "",  # To be filled by student
                "",  # Signature line
            ])

        t = Table(table_data, colWidths=[25, 65, 85, 140, 50, 85, 95])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (3, 1), (3, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#64748b")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))

        elements.append(t)
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Invigilator Signature: ___________________________    Date: _____________", sub_style))
        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def render_duty_roster_pdf(contract: DutyRosterReportContract) -> BytesIO:
        """Generates an Invigilation Duty Roster PDF."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=15, leading=18, alignment=1)
        sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=10, leading=13, alignment=1)

        elements = [
            Paragraph(contract.institution_name, title_style),
            Paragraph(f"FACULTY INVIGILATION DUTY ROSTER — {contract.exam_code}", sub_style),
            Paragraph(
                f"Exam: {contract.exam_title} | Date: {contract.exam_date} ({contract.session_name}) | "
                f"Time: {contract.start_time} to {contract.end_time}",
                sub_style,
            ),
            Spacer(1, 15),
        ]

        table_data = [["Sr", "Faculty Name", "Emp ID", "Dept", "Designation", "Room", "Role", "Signature"]]
        for d in contract.duties:
            table_data.append([
                str(d.sr_no),
                d.faculty_name,
                d.employee_id,
                d.department_code,
                d.designation,
                f"{d.room_number} ({d.building})",
                d.duty_role,
                "",
            ])

        t = Table(table_data, colWidths=[25, 120, 60, 45, 75, 95, 60, 65])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (1, 1), (1, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))

        elements.append(t)
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Controller of Examinations: ___________________________", sub_style))
        doc.build(elements)
        buffer.seek(0)
        return buffer

    # =========================================================================
    # EXCEL RENDERING ENGINES (openpyxl)
    # =========================================================================

    @staticmethod
    def render_seating_chart_excel(contract: SeatingChartReportContract) -> BytesIO:
        """Generates an Excel workbook for the Master Seating Arrangement."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Seating Chart"

        # Headers
        ws.append([contract.institution_name])
        ws.append([f"Master Seating Chart: {contract.plan_code} - {contract.exam_code}"])
        ws.append([f"Subject: {contract.subject_name} | Date: {contract.exam_date} | Session: {contract.session_name}"])
        ws.append([])  # blank row

        col_headers = ["Sr No", "Roll Number", "Student Name", "Department", "Semester", "Room", "Seat Number", "Row", "Col"]
        ws.append(col_headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

        for col_idx in range(1, len(col_headers) + 1):
            cell = ws.cell(row=5, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        for idx, item in enumerate(contract.allocations, start=1):
            row_num = 5 + idx
            row_data = [
                idx,
                item.student_roll,
                item.student_name,
                item.department_code,
                item.semester,
                item.room_number,
                item.seat_number,
                item.row_num,
                item.col_num,
            ]
            ws.append(row_data)
            for c_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num, column=c_idx)
                cell.border = thin_border
                if c_idx not in (3,):
                    cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def render_attendance_sheet_excel(contract: AttendanceSheetReportContract) -> BytesIO:
        """Generates an Excel workbook for the Hall Attendance Sheet."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Sheet"

        ws.append([contract.institution_name])
        ws.append([f"Hall Attendance Sheet — Room {contract.room_number} — Exam {contract.exam_code}"])
        ws.append([f"Subject: {contract.subject_name} | Date: {contract.exam_date} | Session: {contract.session_name}"])
        ws.append([])

        col_headers = ["Sr No", "Roll Number", "Enrollment Number", "Student Name", "Seat Number", "Answer Book No", "Signature"]
        ws.append(col_headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

        for col_idx in range(1, len(col_headers) + 1):
            cell = ws.cell(row=5, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, item in enumerate(contract.candidates, start=1):
            row_data = [idx, item.roll_number, item.enrollment_number, item.student_name, item.seat_number, "", ""]
            ws.append(row_data)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def render_duty_roster_excel(contract: DutyRosterReportContract) -> BytesIO:
        """Generates an Excel workbook for the Invigilation Duty Roster."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Duty Roster"

        ws.append([contract.institution_name])
        ws.append([f"Invigilation Duty Roster — Exam: {contract.exam_code}"])
        ws.append([f"Date: {contract.exam_date} ({contract.session_name}) | Time: {contract.start_time} - {contract.end_time}"])
        ws.append([])

        col_headers = ["Sr No", "Faculty Name", "Employee ID", "Department", "Designation", "Room", "Building", "Role", "Signature"]
        ws.append(col_headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

        for col_idx in range(1, len(col_headers) + 1):
            cell = ws.cell(row=5, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, item in enumerate(contract.duties, start=1):
            ws.append([
                idx,
                item.faculty_name,
                item.employee_id,
                item.department_code,
                item.designation,
                item.room_number,
                item.building,
                item.duty_role,
                "",
            ])

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # =========================================================================
    # GATEWAY DISPATCHERS (Invoked by Flask routes)
    # =========================================================================

    @staticmethod
    def generate_seating_chart(
        plan_id: int,
        format_type: str = "pdf",
        operator_user_id: Optional[int] = None,
        ip_address: Optional[str] = None,
    ) -> Tuple[BytesIO, str, str]:
        """
        Gateway method generating the Master Seating Arrangement Chart in PDF or Excel.
        Returns: (file_stream, filename, mimetype)
        """
        contract = ReportService.build_seating_chart_contract(plan_id)
        fmt = format_type.strip().lower()

        if fmt == "pdf":
            buf = ReportService.render_seating_chart_pdf(contract)
            filename = f"Seating_Chart_{contract.plan_code}.pdf"
            mime_type = "application/pdf"
        elif fmt in ("excel", "xlsx"):
            buf = ReportService.render_seating_chart_excel(contract)
            filename = f"Seating_Chart_{contract.plan_code}.xlsx"
            mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            raise BadRequestError(f"Unsupported report format '{format_type}'. Expected 'pdf' or 'excel'.")

        log_audit(
            action="REPORT_GENERATED",
            entity_type="SeatingPlan",
            entity_id=str(plan_id),
            user_id=operator_user_id,
            details=f"Master seating chart exported in format '{fmt}'.",
            ip_address=ip_address,
        )

        return buf, filename, mime_type

    @staticmethod
    def generate_room_notice(
        plan_id: int,
        room_id: int,
        format_type: str = "pdf",
        operator_user_id: Optional[int] = None,
        ip_address: Optional[str] = None,
    ) -> Tuple[BytesIO, str, str]:
        """
        Gateway method generating the Room Door Notice / Noticeboard Seating List.
        Returns: (file_stream, filename, mimetype)
        """
        contract = ReportService.build_room_notice_contract(plan_id, room_id)
        fmt = format_type.strip().lower()

        if fmt == "pdf":
            buf = ReportService.render_room_notice_pdf(contract)
            filename = f"Room_Notice_{contract.room_number}_{contract.exam_code}.pdf"
            mime_type = "application/pdf"
        else:
            raise BadRequestError("Room notices are only exportable in 'pdf' format.")

        log_audit(
            action="REPORT_GENERATED",
            entity_type="RoomNotice",
            entity_id=f"Plan:{plan_id}_Room:{room_id}",
            user_id=operator_user_id,
            details=f"Room door notice for Room {contract.room_number} exported as PDF.",
            ip_address=ip_address,
        )

        return buf, filename, mime_type

    @staticmethod
    def generate_attendance_sheet(
        plan_id: int,
        room_id: int,
        format_type: str = "pdf",
        operator_user_id: Optional[int] = None,
        ip_address: Optional[str] = None,
    ) -> Tuple[BytesIO, str, str]:
        """
        Gateway method generating the Hall Attendance & Signature Sheet in PDF or Excel.
        Returns: (file_stream, filename, mimetype)
        """
        contract = ReportService.build_attendance_sheet_contract(plan_id, room_id)
        fmt = format_type.strip().lower()

        if fmt == "pdf":
            buf = ReportService.render_attendance_sheet_pdf(contract)
            filename = f"Attendance_Sheet_Room_{contract.room_number}_{contract.exam_code}.pdf"
            mime_type = "application/pdf"
        elif fmt in ("excel", "xlsx"):
            buf = ReportService.render_attendance_sheet_excel(contract)
            filename = f"Attendance_Sheet_Room_{contract.room_number}_{contract.exam_code}.xlsx"
            mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            raise BadRequestError(f"Unsupported report format '{format_type}'. Expected 'pdf' or 'excel'.")

        log_audit(
            action="REPORT_GENERATED",
            entity_type="AttendanceSheet",
            entity_id=f"Plan:{plan_id}_Room:{room_id}",
            user_id=operator_user_id,
            details=f"Attendance sheet for Room {contract.room_number} exported in format '{fmt}'.",
            ip_address=ip_address,
        )

        return buf, filename, mime_type

    @staticmethod
    def generate_duty_roster(
        exam_id: int,
        format_type: str = "pdf",
        operator_user_id: Optional[int] = None,
        ip_address: Optional[str] = None,
    ) -> Tuple[BytesIO, str, str]:
        """
        Gateway method generating the Faculty Invigilation Duty Roster in PDF or Excel.
        Returns: (file_stream, filename, mimetype)
        """
        contract = ReportService.build_duty_roster_contract(exam_id)
        fmt = format_type.strip().lower()

        if fmt == "pdf":
            buf = ReportService.render_duty_roster_pdf(contract)
            filename = f"Duty_Roster_{contract.exam_code}.pdf"
            mime_type = "application/pdf"
        elif fmt in ("excel", "xlsx"):
            buf = ReportService.render_duty_roster_excel(contract)
            filename = f"Duty_Roster_{contract.exam_code}.xlsx"
            mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            raise BadRequestError(f"Unsupported report format '{format_type}'. Expected 'pdf' or 'excel'.")

        log_audit(
            action="REPORT_GENERATED",
            entity_type="DutyRoster",
            entity_id=str(exam_id),
            user_id=operator_user_id,
            details=f"Invigilation duty roster for Exam {contract.exam_code} exported in format '{fmt}'.",
            ip_address=ip_address,
        )

        return buf, filename, mime_type

    @staticmethod
    def get_contracts_meta() -> Dict[str, Any]:
        """Returns JSON schema metadata for all Member 3 <-> Member 4 data contracts."""
        return {
            "version": "1.0-polytechnic",
            "contracts": {
                "SeatingAlgorithmInput": {
                    "description": "Payload sent to Member 4 seating algorithm",
                    "fields": ["exam_id", "exam_code", "students", "rooms", "options"],
                },
                "SeatingAlgorithmOutput": {
                    "description": "Payload returned by Member 4 seating algorithm",
                    "fields": ["allocations", "unallocated_student_ids", "total_allocated", "status"],
                },
                "InvigilationAlgorithmInput": {
                    "description": "Payload sent to Member 4 invigilation scheduler",
                    "fields": ["exam_id", "exam_code", "exam_date", "session_name", "rooms", "teachers", "options"],
                },
                "InvigilationAlgorithmOutput": {
                    "description": "Payload returned by Member 4 scheduler",
                    "fields": ["duties", "unassigned_room_ids", "status"],
                },
                "ReportContracts": [
                    "SeatingChartReportContract",
                    "RoomNoticeReportContract",
                    "AttendanceSheetReportContract",
                    "DutyRosterReportContract",
                ],
            },
        }
